"""Fabricates the demo dataset: one RFx, five vendor replies in five shapes, and the true answers.

Run:  uv run python scripts/generate_dataset.py

Everything is derived from one price matrix, so data/seed/ground_truth.json always matches the
files the vendors "sent". The app never reads ground_truth.json; it exists only to score extraction.
"""
import json
import random
from pathlib import Path

import numpy as np
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageDraw, ImageEnhance, ImageFilter, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

SEED = Path("data/seed")
RESP = SEED / "responses"
FX_USD_INR = 88.0
rng = random.Random(7)

# (description, ply, L, W, H in mm, annual quantity in pieces)
ITEMS = [
    ("E-com shipper carton, plain", 3, 200, 150, 100, 160000),
    ("E-com shipper carton, plain", 3, 250, 200, 150, 130000),
    ("E-com shipper carton, plain", 3, 300, 200, 150, 120000),
    ("E-com shipper carton, plain", 3, 300, 250, 200, 105000),
    ("E-com shipper carton, plain", 3, 350, 250, 200, 90000),
    ("E-com shipper carton, plain", 3, 400, 300, 200, 80000),
    ("Die-cut mailer box, 1-colour print", 3, 220, 160, 60, 145000),
    ("Die-cut mailer box, 1-colour print", 3, 300, 220, 80, 120000),
    ("Inner carton, plain", 3, 180, 120, 90, 200000),
    ("Inner carton, plain", 3, 240, 160, 110, 160000),
    ("Display tray, plain", 3, 400, 300, 80, 65000),
    ("Carton with 12-cell partition", 3, 380, 280, 180, 52000),
    ("Master carton, plain", 5, 400, 300, 300, 80000),
    ("Master carton, plain", 5, 450, 350, 300, 65000),
    ("Master carton, plain", 5, 500, 400, 300, 60000),
    ("Master carton, plain", 5, 500, 400, 400, 52000),
    ("Master carton, plain", 5, 600, 400, 400, 46000),
    ("Master carton, plain", 5, 600, 450, 450, 40000),
    ("Export carton, moisture-resistant", 5, 550, 350, 350, 33000),
    ("Master carton, 2-colour print", 5, 450, 300, 250, 72000),
    ("Master carton, 2-colour print", 5, 480, 320, 280, 52000),
    ("Bottle shipper with partitions", 5, 420, 320, 260, 60000),
    ("Heavy inner carton, plain", 5, 350, 300, 250, 65000),
    ("Tall carton, plain", 5, 400, 400, 600, 26000),
    ("Heavy-duty carton, plain", 7, 600, 500, 500, 16000),
    ("Heavy-duty carton, plain", 7, 700, 500, 500, 13000),
    ("Heavy-duty carton, plain", 7, 800, 600, 500, 10000),
    ("Pallet box, plain", 7, 1000, 800, 600, 5000),
    ("Pallet box, plain", 7, 1200, 800, 800, 4000),
    ("Export heavy-duty carton, moisture-resistant", 7, 750, 550, 550, 8000),
]
BOARD_GSM = {3: 450, 5: 750, 7: 1050}  # combined board weight, g per square metre
RATE_PER_KG = {3: 38.0, 5: 42.0, 7: 46.0}

QUESTIONNAIRE = [
    {"id": "Q1", "question": "Is your plant ISO 9001 certified? Attach a valid certificate.", "mandatory": True,
     "requires_document": "iso_9001_certificate"},
    {"id": "Q2", "question": "Will you supply BCT and burst-strength test reports with every lot?", "mandatory": True},
    {"id": "Q3", "question": "Do you hold FSC chain-of-custody certification?", "mandatory": False},
    {"id": "Q4", "question": "What is your installed capacity in tonnes per month?", "mandatory": False},
    {"id": "Q5", "question": "What is your delivery lead time from purchase order, in days?", "mandatory": False},
    {"id": "Q6", "question": "Do you accept payment at 60 days from invoice? If not, state your terms.", "mandatory": False},
    {"id": "Q7", "question": "State your latest annual turnover and the year you were established.", "mandatory": False},
    {"id": "Q8", "question": "What is your current credit rating, if rated? Name the agency.", "mandatory": False},
]


def build_lines():
    lines = []
    for n, (desc, ply, l, w, h, qty) in enumerate(ITEMS, start=1):
        area_m2 = (2 * (l + w) + 40) * (h + w) / 1e6  # blank size of a regular slotted carton
        weight = round(area_m2 * BOARD_GSM[ply] / 1000, 3)
        base = weight * RATE_PER_KG[ply] * rng.uniform(0.97, 1.08)  # printing, partitions, coating
        lines.append({
            "line_no": n, "item_code": f"CB-{ply}P-{n:03d}", "description": desc, "ply": ply,
            "length_mm": l, "width_mm": w, "height_mm": h, "box_weight_kg": weight,
            "annual_qty": qty, "uom": "piece", "last_year_price_inr": round(base * 0.96, 2),
            "_base": base,
        })
    return lines


def size(line):
    return f"{line['length_mm']} x {line['width_mm']} x {line['height_mm']}"


def write_email(folder, sender, address, subject, received, body, attachments):
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "email.json").write_text(json.dumps({
        "from_name": sender, "from_email": address, "subject": subject,
        "received_at": received, "body": body, "attachments": attachments,
    }, indent=2, ensure_ascii=False))


def iso_certificate(path, company, cert_no, issued, valid_until):
    styles = getSampleStyleSheet()
    centre = ParagraphStyle("c", parent=styles["Normal"], alignment=1, fontSize=12, leading=20)
    doc = SimpleDocTemplate(str(path), pagesize=A4, topMargin=90)
    doc.build([
        Paragraph("CERTIFICATE OF REGISTRATION", ParagraphStyle("t", parent=styles["Title"], fontSize=22)),
        Spacer(1, 24),
        Paragraph("This is to certify that the Quality Management System of", centre),
        Paragraph(f"<b>{company}</b>", ParagraphStyle("n", parent=centre, fontSize=16, leading=28)),
        Paragraph("has been assessed and found to conform to the requirements of", centre),
        Paragraph("<b>ISO 9001:2015</b>", ParagraphStyle("s", parent=centre, fontSize=18, leading=30)),
        Paragraph("Scope: Manufacture and supply of corrugated boxes and cartons.", centre),
        Spacer(1, 30),
        Paragraph(f"Certificate No: {cert_no}", centre),
        Paragraph(f"Date of issue: {issued}", centre),
        Paragraph(f"Valid until: {valid_until}", centre),
        Spacer(1, 40),
        Paragraph("Issued by TransCert Assessment Services Pvt. Ltd.", centre),
    ])


# --- Vendor A: a tidy Excel that ignores the template. Own codes, own order, priced per 100. ---
def vendor_a(lines):
    folder = RESP / "shree_balaji"
    truth = {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Quotation"
    ws.merge_cells("A1:G1")
    ws["A1"] = "SHREE BALAJI PACKAGING PVT. LTD."
    ws["A1"].font = Font(bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:G2")
    ws["A2"] = "Plot 41, MIDC Chakan, Pune 410501  |  Quotation SBP/Q/2026/0914 dated 01-Oct-2026"
    ws.merge_cells("A3:G3")
    ws["A3"] = "Rates are delivered to your Bhiwandi warehouse. GST @ 18% extra. Validity 60 days."
    header = ["S.No", "Our Code", "Particulars", "Size L x W x H (mm)", "Ply", "Rate per 100 Nos (Rs.)", "MOQ (Nos)"]
    for col, text in enumerate(header, start=1):
        cell = ws.cell(row=5, column=col, value=text)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    shuffled = sorted(lines, key=lambda x: (-x["ply"], x["length_mm"] * x["width_mm"] * x["height_mm"]))
    for i, line in enumerate(shuffled, start=1):
        per_100 = round(line["_base"] * 1.00 * rng.uniform(0.96, 1.04) * 100 * 2) / 2
        kind = "RSC" if "carton" in line["description"].lower() else line["description"].split(",")[0]
        extra = line["description"].split(",")[1].strip() if "," in line["description"] else ""
        ws.append([i, f"SBP-{line['ply']}P-{1000 + i * 7}", f"{line['ply']} Ply {kind} {extra}".strip(),
                   size(line), line["ply"], per_100, 1000 if line["ply"] < 7 else 250])
        truth[line["line_no"]] = {"quoted_value": per_100, "quoted_unit": "per 100 pieces", "currency": "INR",
                                  "inr_per_piece": round(per_100 / 100, 4), "status": "quoted"}
    for col, width in zip("ABCDEFG", [6, 16, 44, 22, 6, 24, 12]):
        ws.column_dimensions[col].width = width
    info = wb.create_sheet("Company Profile")
    for row in [
        ("Company", "Shree Balaji Packaging Pvt. Ltd."),
        ("Quality certification", "ISO 9001:2015, certificate attached, valid till March 2028"),
        ("Testing", "BCT and bursting strength report shared with each dispatch lot"),
        ("FSC", "Not held at present"),
        ("Plant capacity", "1,800 MT per month"),
        ("Delivery", "10 days from PO"),
        ("Payment", "45 days from invoice. 60 days not acceptable."),
        ("Turnover (FY 2025-26)", "Rs. 85 crore"),
        ("Established", "2009"),
        ("Credit rating", "CRISIL BBB+ / Stable"),
    ]:
        info.append(row)
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 70
    folder.mkdir(parents=True, exist_ok=True)
    wb.save(folder / "SBP_Quotation_0914.xlsx")
    iso_certificate(folder / "ISO9001_Certificate.pdf", "Shree Balaji Packaging Pvt. Ltd.",
                    "TC/QMS/25/3381", "18 March 2025", "17 March 2028")
    write_email(folder, "Rakesh Agarwal", "sales@shreebalajipack.example", "RE: RFQ corrugated boxes FY27 - our offer",
                "2026-10-01T11:42:00+05:30",
                "Dear Sir,\n\nPlease find attached our quotation and ISO certificate. Company details are in the "
                "second sheet.\n\nRegards,\nRakesh Agarwal\nShree Balaji Packaging",
                ["SBP_Quotation_0914.xlsx", "ISO9001_Certificate.pdf"])
    return {"vendor": "Shree Balaji Packaging Pvt. Ltd.", "format": "Excel, own layout", "freight": "included",
            "lines": truth,
            "questionnaire": {"Q1": "Yes, valid to 17 Mar 2028", "Q2": "Yes", "Q3": "No", "Q4": "1800",
                              "Q5": "10", "Q6": "No, 45 days"},
            "clears_quality_gate": True}


# --- Vendor B: PDF on letterhead, 27 of 30 lines, rebate hidden in a footnote. ---
def vendor_b(lines):
    folder = RESP / "kraftline"
    folder.mkdir(parents=True, exist_ok=True)
    skipped = {12, 28, 29}
    truth = {}
    styles = getSampleStyleSheet()
    small = ParagraphStyle("small", parent=styles["Normal"], fontSize=6.5, leading=8, textColor=colors.grey)
    rows = [["Sr", "Your Item Code", "Description", "Size (mm)", "Rate / pc (Rs.)*"]]
    sr = 0
    for line in lines:
        if line["line_no"] in skipped:
            truth[line["line_no"]] = {"status": "not_quoted", "inr_per_piece": None}
            continue
        sr += 1
        rate = round(line["_base"] * 1.035 * rng.uniform(0.96, 1.04), 2)
        rows.append([sr, line["item_code"], f"{line['ply']}-ply {line['description']}", size(line), f"{rate:.2f}"])
        truth[line["line_no"]] = {"quoted_value": rate, "quoted_unit": "per piece", "currency": "INR",
                                  "inr_per_piece": rate, "inr_per_piece_after_rebate": round(rate * 0.96, 4),
                                  "status": "quoted"}
    table = Table(rows, colWidths=[28, 80, 210, 90, 80], repeatRows=1)
    table.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 8), ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#333333")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ALIGN", (-1, 0), (-1, -1), "RIGHT"),
    ]))
    answers = [
        "1. Quality system: We are an ISO 9001:2015 certified unit. Certificate enclosed (valid up to 09 August 2027).",
        "2. Test reports: BCT and burst factor certificates accompany every lot as standard practice.",
        "3. FSC: Yes, FSC-C104412 chain of custody.",
        "4. Capacity: 2,500 tonnes per month across two corrugators.",
        "5. Lead time: 14 days from receipt of purchase order.",
        "6. Payment: 60 days from invoice date is acceptable.",
        "7. Turnover and standing: Rs. 240 crore in FY 2025-26. In business since 1994.",
        "8. Credit rating: CRISIL A- / Stable.",
    ]
    doc = SimpleDocTemplate(str(folder / "Kraftline_Offer_KL-2609.pdf"), pagesize=A4, topMargin=40, bottomMargin=40)
    doc.build([
        Paragraph("<b>KRAFTLINE INDUSTRIES LIMITED</b>", ParagraphStyle("lh", parent=styles["Title"], fontSize=18)),
        Paragraph("Survey No. 212, Vapi Industrial Estate, Gujarat 396195 &nbsp;|&nbsp; GSTIN 24AAACK4521L1Z6",
                  ParagraphStyle("addr", parent=styles["Normal"], alignment=1, fontSize=8)),
        Spacer(1, 14),
        Paragraph("Offer No. KL/2609/26 &nbsp;&nbsp; Date: 03 October 2026", styles["Normal"]),
        Paragraph("Sub: Your RFQ for corrugated boxes, annual rate contract FY 2026-27", styles["Normal"]),
        Spacer(1, 6),
        Paragraph("We thank you for your enquiry and are pleased to offer as under. Rates are delivered Bhiwandi, "
                  "GST extra as applicable. We regret our inability to offer for partition cartons and pallet "
                  "boxes at this time.", styles["Normal"]),
        Spacer(1, 8), table, Spacer(1, 6),
        Paragraph("* Rates shown are list rates. A rebate of 4% on all items will be passed on by quarterly credit "
                  "note, provided the annual offtake under this contract exceeds Rs. 75,00,000 (basic value). "
                  "Offer valid for 45 days.", small),
        PageBreak(),
        Paragraph("<b>Annexure A: Response to supplier questionnaire</b>", styles["Heading3"]),
        *[Paragraph(a, styles["Normal"]) for a in answers],
    ])
    iso_certificate(folder / "Kraftline_ISO_9001.pdf", "Kraftline Industries Limited",
                    "TC/QMS/24/1907", "10 August 2024", "09 August 2027")
    write_email(folder, "Meera Shah", "meera.shah@kraftline.example", "Kraftline offer KL/2609/26 against your RFQ",
                "2026-10-03T16:05:00+05:30",
                "Dear Team,\n\nOur offer and ISO certificate are attached for your kind consideration.\n\n"
                "Warm regards,\nMeera Shah\nKey Accounts, Kraftline Industries",
                ["Kraftline_Offer_KL-2609.pdf", "Kraftline_ISO_9001.pdf"])
    return {"vendor": "Kraftline Industries Limited", "format": "PDF on letterhead", "freight": "included",
            "conditional_discount": {"pct": 4, "condition": "annual offtake above Rs. 75,00,000"},
            "lines": truth,
            "questionnaire": {"Q1": "Yes, valid to 09 Aug 2027", "Q2": "Yes", "Q3": "Yes", "Q4": "2500",
                              "Q5": "14", "Q6": "Yes"},
            "clears_quality_gate": True}


# --- Vendor C: Word document, commercials written as a paragraph, quoted in USD per kg. ---
def vendor_c(lines):
    folder = RESP / "pacificcorr"
    folder.mkdir(parents=True, exist_ok=True)
    usd_per_kg = {3: 0.42, 5: 0.465, 7: 0.51}
    pallet_rate = 0.55
    truth = {}
    for line in lines:
        rate = pallet_rate if line["line_no"] in (28, 29) else usd_per_kg[line["ply"]]
        truth[line["line_no"]] = {"quoted_value": rate, "quoted_unit": "per kg", "currency": "USD",
                                  "inr_per_piece": round(rate * line["box_weight_kg"] * FX_USD_INR, 4),
                                  "status": "quoted"}
    doc = Document()
    doc.add_heading("PacificCorr Global Pte. Ltd.", level=1)
    doc.add_paragraph("Commercial Proposal: Corrugated Packaging Supply, FY 2026-27\n06 October 2026")
    doc.add_heading("Our understanding", level=2)
    doc.add_paragraph("Thank you for inviting PacificCorr to participate. We have reviewed all thirty items in your "
                      "enquiry and confirm we can supply the full range from our Johor Bahru plant.")
    doc.add_heading("Commercial offer", level=2)
    doc.add_paragraph(
        "Because kraft liner is traded in dollars, our offer is on a weight basis in US dollars rather than per box. "
        "For all 3-ply items we offer USD 0.42 per kg, and for all 5-ply items USD 0.465 per kg. 7-ply heavy-duty "
        "items are offered at USD 0.51 per kg, with the exception of the two pallet boxes (your lines 28 and 29), "
        "which need a wider machine and are offered at USD 0.55 per kg. Weights will be billed as per the box "
        "weights stated in your specification. Prices are CIF Nhava Sheva; Indian customs duty, port clearance and "
        "inland freight to your warehouse are to your account. This offer is valid for 15 days from the date above. "
        "Payment is by irrevocable letter of credit at sight.")
    doc.add_heading("Quality and capability", level=2)
    doc.add_paragraph(
        "Our plant is certified to ISO 9001:2015 (certificate attached, valid through 30 November 2027) and holds "
        "FSC chain of custody. A burst and compression test report is issued for each production lot and travels "
        "with the shipping documents. Installed capacity is 6,000 tonnes a month. Please allow 35 days from order "
        "to arrival at port. We are unable to offer open credit to new customers. Group revenue was USD 60 million in "
        "2025 and we have been operating since 2001. We are not rated by an Indian credit rating agency.")
    doc.save(str(folder / "PacificCorr_Proposal.docx"))
    iso_certificate(folder / "PCG_ISO9001.pdf", "PacificCorr Global Pte. Ltd.",
                    "TC/QMS/24/2750", "01 December 2024", "30 November 2027")
    write_email(folder, "Daniel Lim", "daniel.lim@pacificcorr.example", "PacificCorr proposal for your corrugated RFQ",
                "2026-10-06T09:18:00+05:30",
                "Hello,\n\nAttached is our proposal with ISO certificate. Happy to get on a call to walk you "
                "through it.\n\nBest,\nDaniel Lim\nRegional Sales Director",
                ["PacificCorr_Proposal.docx", "PCG_ISO9001.pdf"])
    return {"vendor": "PacificCorr Global Pte. Ltd.", "format": "Word, commercials in a paragraph",
            "freight": "CIF port only; duty and inland freight excluded", "fx_usd_inr": FX_USD_INR, "lines": truth,
            "questionnaire": {"Q1": "Yes, valid to 30 Nov 2027", "Q2": "Yes", "Q3": "Yes", "Q4": "6000",
                              "Q5": "35", "Q6": "No, LC at sight"},
            "clears_quality_gate": True}


# --- Vendor D: phone photo of a printed rate card, taken at an angle. Cheapest, but ISO has lapsed. ---
def perspective_coeffs(dst, src):
    rows = []
    for (x, y), (u, v) in zip(dst, src):
        rows.append([x, y, 1, 0, 0, 0, -u * x, -u * y])
        rows.append([0, 0, 0, x, y, 1, -v * x, -v * y])
    return np.linalg.solve(np.array(rows, dtype=float), np.array(src, dtype=float).reshape(8))


def vendor_d(lines):
    folder = RESP / "maruti"
    folder.mkdir(parents=True, exist_ok=True)
    truth = {}
    w, h = 1400, 2060
    card = Image.new("RGB", (w, h), "#fbfaf4")
    draw = ImageDraw.Draw(card)
    big, mid, body = (ImageFont.load_default(size=s) for s in (54, 30, 31))
    draw.text((w // 2, 60), "MARUTI CORRUGATORS", font=big, fill="#7a1010", anchor="mm")
    draw.text((w // 2, 118), "Gala 7, Vasai East, Palghar  |  RATE CARD  |  w.e.f. 1 Oct 2026", font=mid,
              fill="black", anchor="mm")
    cols = [60, 180, 760, 1000, 1340]
    top, row_h = 170, 56
    for x, text in zip(cols, ["No.", "Item", "Size (mm)", "Ply", "Rate Rs./pc"]):
        draw.text((x, top + 10), text, font=mid, fill="black", anchor="ra" if text.startswith("Rate") else "la")
    draw.line([(40, top + row_h - 4), (w - 40, top + row_h - 4)], fill="black", width=3)
    smudged_line = 17
    for i, line in enumerate(lines):
        y = top + row_h * (i + 1) + 8
        rate = round(line["_base"] * 0.945 * rng.uniform(0.96, 1.04) * 20) / 20
        draw.text((cols[0], y), str(line["line_no"]), font=body, fill="black")
        draw.text((cols[1], y), line["description"].split(",")[0][:30], font=body, fill="black")
        draw.text((cols[2], y), size(line).replace(" ", ""), font=body, fill="black")
        draw.text((cols[3], y), str(line["ply"]), font=body, fill="black")
        draw.text((cols[4], y), f"{rate:.2f}", font=body, fill="black", anchor="ra")
        draw.line([(40, y + row_h - 12), (w - 40, y + row_h - 12)], fill="#bbbbbb", width=1)
        if line["line_no"] == smudged_line:  # a thumb smear across the rate
            smear = Image.new("RGBA", card.size, (0, 0, 0, 0))
            ImageDraw.Draw(smear).ellipse([cols[4] - 95, y - 2, cols[4] - 30, y + 40], fill=(70, 60, 50, 150))
            card = Image.alpha_composite(card.convert("RGBA"), smear.filter(ImageFilter.GaussianBlur(5))).convert("RGB")
            draw = ImageDraw.Draw(card)
        truth[line["line_no"]] = {"quoted_value": rate, "quoted_unit": "per piece", "currency": "INR",
                                  "inr_per_piece": rate, "status": "quoted"}
    draw.text((60, h - 150), "Rates ex-works Vasai. Transport extra at actuals. GST 18% extra.", font=mid, fill="black")
    draw.text((60, h - 105), "Payment: 30 days. Delivery within 7 days of order.", font=mid, fill="black")
    draw.text((60, h - 60), "ISO 9001:2015 certified company", font=mid, fill="#7a1010")

    out_w, out_h = 1700, 2300
    desk = Image.new("RGB", (out_w, out_h), "#6b5138")
    corners = [(190, 150), (1560, 80), (1640, 2210), (70, 2140)]  # where the card lands in the photo
    coeffs = perspective_coeffs(corners, [(0, 0), (w, 0), (w, h), (0, h)])
    warped = card.transform((out_w, out_h), Image.PERSPECTIVE, list(coeffs), Image.BICUBIC)
    mask = Image.new("L", (w, h), 255).transform((out_w, out_h), Image.PERSPECTIVE, list(coeffs))
    desk.paste(warped, (0, 0), mask)
    shade = Image.linear_gradient("L").resize((out_w, out_h)).rotate(35)  # uneven room lighting
    photo = Image.composite(desk, ImageEnhance.Brightness(desk).enhance(0.72), shade)
    photo = photo.rotate(-3.5, resample=Image.BICUBIC, fillcolor="#6b5138").filter(ImageFilter.GaussianBlur(0.9))
    photo.save(folder / "IMG_20261005_184212.jpg", quality=68)

    iso_certificate(folder / "ISO_cert.pdf", "Maruti Corrugators", "TC/QMS/23/0462", "01 July 2023", "30 June 2026")
    write_email(folder, "Jignesh Patel", "maruticorrugators@mail.example", "Rate", "2026-10-05T18:47:00+05:30",
                "Sir rate card photo attached. ISO certificate also attached. Test report we give with every lot "
                "no problem. Capacity 600 ton monthly. 60 days payment not possible sir, 30 days. Turnover around 6 crore, "
                "unit running since 2016.\n\n"
                "Jignesh Patel\nMaruti Corrugators\nSent from my phone",
                ["IMG_20261005_184212.jpg", "ISO_cert.pdf"])
    return {"vendor": "Maruti Corrugators", "format": "Phone photo of a printed rate card",
            "freight": "excluded (ex-works, transport at actuals)", "expected_low_confidence_lines": [smudged_line],
            "lines": truth,
            "questionnaire": {"Q1": "Claims yes, but the attached certificate expired 30 Jun 2026", "Q2": "Yes",
                              "Q3": "Not answered", "Q4": "600", "Q5": "7", "Q6": "No, 30 days"},
            "clears_quality_gate": False}


# --- Vendor E: a one-line email. Two rates per kg, "rest same as last year", freight extra. ---
def vendor_e(lines):
    folder = RESP / "om_sai"
    truth = {}
    for line in lines:
        if line["ply"] in (3, 5):
            rate = {3: 38.0, 5: 42.0}[line["ply"]]
            truth[line["line_no"]] = {"quoted_value": rate, "quoted_unit": "per kg", "currency": "INR",
                                      "inr_per_piece": round(rate * line["box_weight_kg"], 4), "status": "quoted"}
        else:
            truth[line["line_no"]] = {"quoted_value": None, "quoted_unit": None, "currency": "INR",
                                      "inr_per_piece": line["last_year_price_inr"],
                                      "status": "assumed_same_as_last_year"}
    write_email(folder, "Suresh Naik", "omsaicartons@mail.example", "Re: RFQ", "2026-10-07T21:10:00+05:30",
                "₹42/kg for the 5-ply, 38 for the 3-ply, rest same as last year, freight extra.\n\n"
                "ISO is there, will send certificate.\n\nSuresh", [])
    return {"vendor": "Om Sai Cartons", "format": "One-line email", "freight": "excluded (amount not stated)",
            "lines": truth,
            "questionnaire": {"Q1": "Claims yes, no certificate provided", "Q2": "Not answered",
                              "Q3": "Not answered", "Q4": "Not answered", "Q5": "Not answered", "Q6": "Not answered"},
            "clears_quality_gate": False}


def main():
    lines = build_lines()
    truth = {"note": "True values behind the vendor files. Used only to score extraction, never read by the app.",
             "fx_usd_inr": FX_USD_INR,
             "vendors": {name: fn(lines) for name, fn in [
                 ("shree_balaji", vendor_a), ("kraftline", vendor_b), ("pacificcorr", vendor_c),
                 ("maruti", vendor_d), ("om_sai", vendor_e)]}}
    public_lines = [{k: v for k, v in line.items() if not k.startswith("_")} for line in lines]
    rfx = {
        "rfx_id": "RFQ-2026-PKG-014",
        "title": "Corrugated boxes, annual rate contract FY 2026-27",
        "buyer": {"company": "Nimbus Home & Personal Care Ltd.", "name": "Pavan Vepada", "role": "Category Buyer, Packaging"},
        "issued_on": "2026-09-28", "due_on": "2026-10-07",
        "scope": "Supply of 30 corrugated box SKUs to the Bhiwandi warehouse for 12 months from 1 November 2026.",
        "terms": {"currency": "INR", "price_basis": "per piece, delivered Bhiwandi, GST extra",
                  "payment": "60 days from invoice", "validity_days": 60},
        "quality_gate": "A vendor clears the quality gate only if every mandatory question is satisfied with evidence.",
        "questionnaire": QUESTIONNAIRE, "line_items": public_lines,
    }
    (SEED / "rfx.json").write_text(json.dumps(rfx, indent=2))
    (SEED / "ground_truth.json").write_text(json.dumps(truth, indent=2))

    last_year = sum(l["last_year_price_inr"] * l["annual_qty"] for l in public_lines)
    print(f"30 lines, last year's spend Rs. {last_year / 1e7:.2f} crore")
    for key, vendor in truth["vendors"].items():
        quoted = [(n, v) for n, v in vendor["lines"].items() if v["inr_per_piece"] is not None]
        total = sum(v["inr_per_piece"] * lines[n - 1]["annual_qty"] for n, v in quoted)
        print(f"  {key:13s} {len(quoted):2d} lines  Rs. {total / 1e7:.2f} crore  gate={vendor['clears_quality_gate']}")


if __name__ == "__main__":
    main()
