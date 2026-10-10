"""HTTP endpoints used by the web app."""
import io
import json
import re
from datetime import datetime
from pathlib import Path

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel

from app import analyst, clarify, comparison, copilot, db
from app.config import settings
from app.extraction.readers import ReadDoc, read_file
from app.extraction.schema import VendorExtraction
from app.pipeline import process_vendor, write_quotes



async def select_event(request: Request):
    """Every request names the RFx it is about. db.connect() then opens that RFx's own database."""
    event = request.headers.get("x-rfx") or request.query_params.get("rfx") or db.DEMO
    if not re.fullmatch(r"[a-z0-9_]{1,40}", event):
        raise HTTPException(400, "Bad RFx id")
    db.current_event.set(event)


router = APIRouter(prefix="/api", dependencies=[Depends(select_event)])


# --- RFx events: the buyer's list of requirements ---

def _event_summary(event: str) -> dict:
    db.current_event.set(event)
    conn = db.connect()
    rfx = db.get_rfx(conn)
    draft = copilot.get_draft(conn)
    replies = conn.execute("SELECT COUNT(*) FROM vendors").fetchone()[0]
    review = conn.execute("SELECT COUNT(*) FROM quotes WHERE needs_review = 1").fetchone()[0]
    return {"key": event, "title": (rfx or draft)["title"] or "Untitled RFx",
            "rfx_id": rfx["rfx_id"] if rfx else None, "lines": len((rfx or draft)["line_items"]), "replies": replies,
            "needs_review": review, "issued_on": rfx.get("issued_on") if rfx else None, "due_on": (rfx or draft).get("due_on"),
            "scope": (rfx or draft).get("scope") or "",
            "status": "draft" if not rfx else "awaiting_replies" if replies == 0 else "comparing"}


@router.get("/events")
def list_events():
    keys = sorted((p.stem for p in db.events_dir().glob("*.db")), key=lambda k: (k != db.DEMO, k))
    return [_event_summary(k) for k in keys]


@router.post("/events")
def new_event():
    event = "rfx_" + datetime.now().strftime("%m%d_%H%M%S")
    db.current_event.set(event)
    db.connect().close()
    return _event_summary(event)


@router.delete("/events/{event}")
def delete_event(event: str):
    if event == db.DEMO or not re.fullmatch(r"[a-z0-9_]{1,40}", event):
        raise HTTPException(403, "The demo RFx cannot be deleted. Use reset instead.")
    db.path(event).unlink(missing_ok=True)
    return {"ok": True}


@router.post("/events/demo/reset")
def reset_demo():
    """Puts the demo RFx back to its freshly extracted state: corrections, uploads and FX changes are undone."""
    db.restore_demo(force=True)
    return {"ok": True}



@router.get("/comparison")
def get_comparison():
    conn = db.connect()
    if db.get_rfx(conn) is None:
        return {"rfx": None, "vendors": [], "lines": []}
    data = comparison.build(conn)
    accuracy = Path("data/accuracy.json")  # true answers exist only for the fabricated demo RFx
    data["accuracy"] = json.loads(accuracy.read_text()) if accuracy.exists() and db.current_event.get() == db.DEMO else None
    return data


@router.get("/comparison/export")
def export_comparison():
    """The whole comparison as a workbook: prices with their status, vendor terms, and every caveat."""
    data = comparison.build(db.connect())
    vendors = data["vendors"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Prices"
    ws.append(["Line", "Item", "Qty / year", "Last year (Rs.)"] + [x for v in vendors for x in (v["name"], "Status")])
    for l in data["lines"]:
        row = [l["line_no"], l["description"], l["annual_qty"], l["last_year_price_inr"]]
        for v in vendors:
            c = v["cells"].get(l["line_no"], {})
            row += [c.get("price_inr"), (c.get("status") or "") + (", needs review" if c.get("needs_review") else "")]
        ws.append(row)
    terms = wb.create_sheet("Vendors")
    terms.append(["Vendor", "Overall risk", "Risk detail", "Quality gate", "Why", "Freight", "Freight note", "Conditional rebate %", "Rebate condition",
                  "Payment", "Lead time (days)", "Validity (days)", "Lines priced", "Total priced (Rs.)"])
    for v in vendors:
        terms.append([v["name"], v["risk"]["overall"] + (" (incomplete information)" if v["risk"]["incomplete"] else ""),
                      "; ".join(f"{d['name']}: {d['level']}. {' '.join(d['reasons'])}" for d in v["risk"]["dimensions"]),
                      v["quality_gate"], v["quality_gate_reason"], v["freight"], v["freight_note"],
                      v["conditional_discount_pct"], v["conditional_discount_condition"], v["payment_terms"],
                      v["lead_time_days"], v["validity_days"], v["summary"]["lines_priced"], round(v["summary"]["total_inr"])])
    notes = wb.create_sheet("Caveats")
    notes.append(["Vendor", "Line", "Status", "How the price was reached", "Why it needs care"])
    for v in vendors:
        for c in v["cells"].values():
            if c["needs_review"] or c["status"] in ("assumed", "unclear", "derived"):
                notes.append([v["name"], c["line_no"], c["status"], c["formula"], " ".join(c["review_reasons"])])
    notes.append([])
    notes.append([f"Prices are rupees per piece before GST. USD converted at {data['fx_usd_inr']:g}. "
                  f"Exported {datetime.now():%d %b %Y %H:%M}."])
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="comparison.xlsx"'})


@router.post("/vendors/{vendor_id}/clarification")
def clarification(vendor_id: str):
    return clarify.draft_email(vendor_id)


@router.get("/vendors/{vendor_id}/documents/{filename}")
def get_document(vendor_id: str, filename: str):
    """Text and reading mode of one document, for the source panel."""
    row = db.connect().execute("SELECT * FROM documents WHERE vendor_id = ? AND filename = ?",
                               (vendor_id, filename)).fetchone()
    if not row:
        raise HTTPException(404, "Document not found")
    return dict(row)


@router.get("/vendors/{vendor_id}/files/{filename}")
def get_file(vendor_id: str, filename: str):
    """The original file the vendor sent."""
    row = db.connect().execute("SELECT folder FROM vendors WHERE vendor_id = ?", (vendor_id,)).fetchone()
    folder = Path(row["folder"]) if row else None
    if not folder or filename not in {p.name for p in folder.iterdir()}:
        raise HTTPException(404, "File not found")
    return FileResponse(folder / filename, filename=filename, content_disposition_type="inline")


class Correction(BaseModel):
    price_inr: float | None  # None means "this vendor has no price for this line"
    note: str = ""


@router.patch("/quotes/{vendor_id}/{line_no}")
def correct_quote(vendor_id: str, line_no: int, body: Correction):
    """The buyer confirms or corrects a cell. The cell is then protected from re-runs and the change is logged."""
    conn = db.connect()
    row = conn.execute("SELECT * FROM quotes WHERE vendor_id = ? AND line_no = ?", (vendor_id, line_no)).fetchone()
    if not row:
        raise HTTPException(404, "Quote not found")
    same = body.price_inr is not None and row["price_inr"] is not None and abs(body.price_inr - row["price_inr"]) < 0.005
    action = "confirmed" if same else "corrected"
    was = "no price" if row["price_inr"] is None else f"Rs. {row['price_inr']:.2f}"
    now = "no price" if body.price_inr is None else f"Rs. {body.price_inr:.2f}"
    detail = f"{vendor_id} line {line_no}: {was} {'confirmed' if same else '-> ' + now}" + (f" ({body.note})" if body.note else "")
    formula = row["formula"] + (" Confirmed by buyer." if same else f" Corrected by buyer from {was} to {now}.")
    price = row["price_inr"] if same else body.price_inr  # a confirmation keeps the exact stored price
    conn.execute("UPDATE quotes SET price_inr = ?, status = ?, needs_review = 0, buyer_override = 1, formula = ? "
                 "WHERE vendor_id = ? AND line_no = ?",
                 (price, row["status"] if price is not None else "missing", formula, vendor_id, line_no))
    db.log(conn, "buyer", action, detail)
    conn.commit()
    return {"ok": True}


class FxRate(BaseModel):
    fx_usd_inr: float


@router.put("/settings/fx")
def set_fx(body: FxRate):
    """Changes the USD rate and recalculates every cell from the stored extraction. No model call."""
    if not 50 <= body.fx_usd_inr <= 150:
        raise HTTPException(422, "Rate must be between 50 and 150")
    conn = db.connect()
    old = db.get_fx(conn)
    conn.execute("UPDATE app_settings SET value = ? WHERE key = 'fx_usd_inr'", (str(body.fx_usd_inr),))
    for v in conn.execute("SELECT vendor_id, extraction FROM vendors").fetchall():
        docs = {d["filename"]: ReadDoc(d["filename"], d["read_mode"], d["read_mode_reason"], d["text"] or "",
                                       ["image"] if "vision" in d["read_mode"] else [])
                for d in conn.execute("SELECT * FROM documents WHERE vendor_id = ?", (v["vendor_id"],))}
        write_quotes(conn, v["vendor_id"], VendorExtraction.model_validate_json(v["extraction"]), docs)
    db.log(conn, "buyer", "fx changed", f"USD rate {old:g} -> {body.fx_usd_inr:g}")
    conn.commit()
    return {"ok": True}


@router.get("/audit")
def audit():
    return [dict(r) for r in db.connect().execute("SELECT * FROM audit_log ORDER BY id DESC LIMIT 200")]


class Question(BaseModel):
    question: str
    chat_id: str | None = None


@router.post("/analyst")
def ask_analyst(body: Question):
    return analyst.ask_in_chat(body.question, body.chat_id)


@router.get("/analyst/chats")
def analyst_chats():
    return analyst.list_chats()


@router.get("/analyst/chats/{chat_id}")
def analyst_chat(chat_id: str):
    chat = analyst.get_chat(chat_id)
    if not chat:
        raise HTTPException(404, "Conversation not found")
    return chat


@router.delete("/analyst/chats/{chat_id}")
def delete_analyst_chat(chat_id: str):
    conn = db.connect()
    conn.execute("DELETE FROM analyst_chats WHERE chat_id = ?", (chat_id,))
    conn.commit()
    return {"ok": True}


class Export(BaseModel):
    title: str
    columns: list[str]
    rows: list[list]
    how: str = ""
    notes: list[str] = []


@router.post("/export")
def export_excel(body: Export):
    """Any table the analyst shows can be downloaded, with a second sheet saying how it was produced."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Result"
    ws.append(body.columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in body.rows:
        ws.append(row)
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width = min(48, max(len(str(c.value or "")) for c in column) + 2)
    about = wb.create_sheet("How this was produced")
    conn = db.connect()
    for line in [body.title, f"Generated {datetime.now():%d %b %Y %H:%M} from {db.get_rfx(conn)['rfx_id']}",
                 f"USD rate used: {db.get_fx(conn):g}", "Prices are rupees per piece before GST.", "",
                 "Method:", body.how, "", "Caveats:", *(body.notes or ["None recorded for this table."])]:
        about.append([line])
    about.column_dimensions["A"].width = 120
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="analysis.xlsx"'})


# --- Vendor replies: add your own, read again, remove ---

UPLOADS = Path(settings.data_dir) / "uploads"


@router.post("/vendors/upload")
def upload_reply(vendor_name: str = Form(""), email_text: str = Form(""), files: list[UploadFile] = File(default=[])):
    """Adds a new vendor reply from any files and runs the real extraction on it."""
    if not files and not email_text.strip():
        raise HTTPException(422, "Add at least one file or some email text.")
    vendor_id = "upload_" + datetime.now().strftime("%H%M%S")
    conn = db.connect()
    if db.get_rfx(conn) is None:
        raise HTTPException(422, "Send the RFx before adding replies to it.")
    folder = UPLOADS / db.current_event.get() / vendor_id
    folder.mkdir(parents=True)
    for upload in files:
        (folder / re.sub(r"[^\w.() -]", "_", Path(upload.filename).name)).write_bytes(upload.file.read())
    (folder / "email.json").write_text(json.dumps({
        "from_name": vendor_name or "Uploaded vendor", "from_email": "uploaded@demo", "subject": "Reply added by buyer",
        "received_at": datetime.now().astimezone().isoformat(timespec="seconds"), "body": email_text,
        "attachments": [u.filename for u in files]}))
    extraction = process_vendor(conn, vendor_id, folder)
    return {"vendor_id": vendor_id, "name": extraction.vendor_name}


@router.post("/vendors/{vendor_id}/reread")
def reread(vendor_id: str):
    """Runs extraction again on the original files. Nothing is cached: this is a fresh model call."""
    conn = db.connect()
    row = conn.execute("SELECT folder FROM vendors WHERE vendor_id = ?", (vendor_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Vendor not found")
    process_vendor(conn, vendor_id, Path(row["folder"]))
    return {"ok": True}


@router.delete("/vendors/{vendor_id}")
def remove_vendor(vendor_id: str):
    if not vendor_id.startswith("upload_"):
        raise HTTPException(403, "Only replies you added can be removed.")
    conn = db.connect()
    for table in ("vendors", "documents", "quotes", "questionnaire"):
        conn.execute(f"DELETE FROM {table} WHERE vendor_id = ?", (vendor_id,))
    db.log(conn, "buyer", "removed reply", vendor_id)
    conn.commit()
    return {"ok": True}


# --- RFx co-pilot ---

class Message(BaseModel):
    message: str


@router.get("/rfx-draft")
def get_draft():
    draft = copilot.get_draft(db.connect())
    return {**draft, "email": copilot.outgoing_email(draft)}


@router.post("/rfx-draft/chat")
def draft_chat(body: Message):
    draft = copilot.chat(body.message)
    return {**draft, "email": copilot.outgoing_email(draft)}


@router.post("/rfx-draft/upload")
def draft_upload(message: str = Form(""), files: list[UploadFile] = File(...)):
    """The buyer hands the co-pilot a requirement document instead of typing it out."""
    folder = UPLOADS / db.current_event.get() / "requirement"
    folder.mkdir(parents=True, exist_ok=True)
    docs = []
    for upload in files:
        path = folder / re.sub(r"[^\w.() -]", "_", Path(upload.filename).name)
        path.write_bytes(upload.file.read())
        docs += read_file(path)
    unreadable = [d.filename for d in docs if d.mode == "unsupported"]
    if unreadable:
        raise HTTPException(422, f"Cannot read {', '.join(unreadable)}. Use Word, PDF, Excel, text or an image.")
    draft = copilot.chat(message, docs)
    return {**draft, "email": copilot.outgoing_email(draft)}


@router.delete("/rfx-draft")
def reset_draft():
    conn = db.connect()
    conn.execute("DELETE FROM app_settings WHERE key = 'rfx_draft'")
    conn.commit()
    return get_draft()


@router.post("/rfx-draft/send")
def send_draft():
    """Stubbed send: no email leaves the system. The draft becomes a live RFx that replies can be added to."""
    conn = db.connect()
    draft = copilot.get_draft(conn)
    if not draft["line_items"]:
        raise HTTPException(422, "Add at least one line item before sending.")
    if db.get_rfx(conn):
        raise HTTPException(409, "This RFx has already been sent.")
    draft["sent_at"] = datetime.now().isoformat(timespec="seconds")
    copilot.save_draft(conn, draft)
    db.load_rfx(conn, {
        "rfx_id": "RFQ-" + db.current_event.get().removeprefix("rfx_").upper().replace("_", "-"),
        "title": draft["title"] or "Untitled RFx", "scope": draft["scope"],
        "buyer": {"company": "Nimbus Home & Personal Care Ltd.", "name": "Pavan Vepada", "role": "Category Buyer, Packaging"},
        "issued_on": draft["sent_at"][:10], "due_on": draft["due_on"], "vendors": draft["vendors"],
        "terms": draft["terms"], "questionnaire": draft["questionnaire"],
        "quality_gate": "A vendor clears the quality gate only if every mandatory question is satisfied with evidence.",
        "line_items": [{**l, "item_code": f"L{l['line_no']:03d}", "box_weight_kg": l.get("unit_weight_kg") or l.get("box_weight_kg"),
                        "last_year_price_inr": None} for l in draft["line_items"]],
    })
    db.log(conn, "buyer", "rfx sent (stubbed)", f"{draft['title']} to {len(draft['vendors'])} vendors")
    conn.commit()
    return {**draft, "email": copilot.outgoing_email(draft)}


@router.get("/rfx-draft/export")
def export_draft():
    draft = copilot.get_draft(db.connect())
    wb = Workbook()
    ws = wb.active
    ws.title = "Line items"
    ws.append(["Line", "Description", "Specification", "Unit", "Unit weight (kg)", "Annual qty", "Your price", "Price unit"])
    for l in draft["line_items"]:
        ws.append([l["line_no"], l["description"], l.get("specification"), l.get("uom"), l.get("unit_weight_kg"), l.get("annual_qty")])
    sheet = wb.create_sheet("Questionnaire")
    sheet.append(["No.", "Question", "Mandatory", "Your answer"])
    for q in draft["questionnaire"]:
        sheet.append([q["id"], q["question"], "Yes" if q["mandatory"] else "No"])
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="RFx_line_items.xlsx"'})
